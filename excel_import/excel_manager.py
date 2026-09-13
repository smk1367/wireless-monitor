"""
Independent Excel import/export module for Wireless Monitor.

IP address is the stable key. Imported values are stored in the existing
manual_* fields and effective fields of the devices table.
"""

import io
import ipaddress
import re
import sqlite3
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("IP", "ip_address"),
    ("Latitude", "latitude"),
    ("Longitude", "longitude"),
    ("Province", "province"),
    ("City", "city"),
    ("Antenna Gain", "antenna_gain"),
    ("Polarization", "polarization"),
    ("Capacity", "capacity"),
]

HEADER_ALIASES = {
    "ip": "IP",
    "ip address": "IP",
    "ip_address": "IP",
    "latitude": "Latitude",
    "lat": "Latitude",
    "longitude": "Longitude",
    "lon": "Longitude",
    "lng": "Longitude",
    "province": "Province",
    "state": "Province",
    "city": "City",
    "antenna gain": "Antenna Gain",
    "antenna_gain": "Antenna Gain",
    "gain": "Antenna Gain",
    "polarization": "Polarization",
    "polarisation": "Polarization",
    "capacity": "Capacity",
}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _normalize_header(value):
    text = _text(value).lower()
    return re.sub(r"[\s_\-]+", " ", text)


def _read_rows(fileobj):
    workbook = openpyxl.load_workbook(
        fileobj,
        read_only=True,
        data_only=True,
    )
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(values_only=True), None)
        if first_row is None:
            raise ValueError("Excel file is empty.")

        positions = {}
        for index, raw in enumerate(first_row):
            canonical = HEADER_ALIASES.get(_normalize_header(raw))
            if canonical and canonical not in positions:
                positions[canonical] = index

        required = [label for label, _ in COLUMNS]
        missing = [label for label in required if label not in positions]
        if missing:
            raise ValueError(
                "Missing required Excel columns: " + ", ".join(missing)
            )

        records = []
        for excel_row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            ip_pos = positions["IP"]
            ip_raw = row[ip_pos] if len(row) > ip_pos else None
            ip = _text(ip_raw)
            if not ip:
                continue
            try:
                ip = str(ipaddress.ip_address(ip))
            except ValueError as exc:
                raise ValueError(
                    f"Invalid IP at Excel row {excel_row_number}: {ip!r}"
                ) from exc

            record = {"ip_address": ip, "excel_row": excel_row_number}
            for label, field in COLUMNS[1:]:
                pos = positions[label]
                record[field] = _text(row[pos]) if len(row) > pos else ""
            records.append(record)

        return records
    finally:
        workbook.close()


def import_excel(fileobj, db_path):
    """Import antenna/site data by IP. Blank cells preserve old values."""
    rows = _read_rows(fileobj)

    seen = set()
    duplicates = []
    for item in rows:
        if item["ip_address"] in seen:
            duplicates.append(item["ip_address"])
        seen.add(item["ip_address"])
    if duplicates:
        raise ValueError(
            "Duplicate IPs in Excel: " + ", ".join(sorted(set(duplicates)))
        )

    conn = sqlite3.connect(db_path, timeout=30)
    try:
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("BEGIN")
        updated = 0
        not_found = []

        for item in rows:
            found = conn.execute(
                "SELECT id FROM devices WHERE ip_address = ?",
                (item["ip_address"],),
            ).fetchone()
            if not found:
                not_found.append(item["ip_address"])
                continue

            conn.execute(
                """
                UPDATE devices
                SET
                    manual_latitude = CASE WHEN ? != '' THEN ? ELSE manual_latitude END,
                    manual_longitude = CASE WHEN ? != '' THEN ? ELSE manual_longitude END,
                    manual_province = CASE WHEN ? != '' THEN ? ELSE manual_province END,
                    manual_city = CASE WHEN ? != '' THEN ? ELSE manual_city END,
                    manual_antenna_gain = CASE WHEN ? != '' THEN ? ELSE manual_antenna_gain END,
                    manual_polarization = CASE WHEN ? != '' THEN ? ELSE manual_polarization END,
                    manual_capacity = CASE WHEN ? != '' THEN ? ELSE manual_capacity END,
                    latitude = CASE WHEN ? != '' THEN ? ELSE latitude END,
                    longitude = CASE WHEN ? != '' THEN ? ELSE longitude END,
                    province = CASE WHEN ? != '' THEN ? ELSE province END,
                    city = CASE WHEN ? != '' THEN ? ELSE city END,
                    antenna_gain = CASE WHEN ? != '' THEN ? ELSE antenna_gain END,
                    polarization = CASE WHEN ? != '' THEN ? ELSE polarization END,
                    capacity = CASE WHEN ? != '' THEN ? ELSE capacity END,
                    last_updated = datetime('now')
                WHERE ip_address = ?
                """,
                (
                    item["latitude"], item["latitude"],
                    item["longitude"], item["longitude"],
                    item["province"], item["province"],
                    item["city"], item["city"],
                    item["antenna_gain"], item["antenna_gain"],
                    item["polarization"], item["polarization"],
                    item["capacity"], item["capacity"],
                    item["latitude"], item["latitude"],
                    item["longitude"], item["longitude"],
                    item["province"], item["province"],
                    item["city"], item["city"],
                    item["antenna_gain"], item["antenna_gain"],
                    item["polarization"], item["polarization"],
                    item["capacity"], item["capacity"],
                    item["ip_address"],
                ),
            )
            updated += 1

        conn.commit()
        return {
            "status": "ok",
            "rows_read": len(rows),
            "updated": updated,
            "not_found": len(not_found),
            "not_found_ips": not_found,
        }
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _new_workbook(sheet_title="Antenna Information"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append([label for label, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    widths = [20, 16, 16, 20, 20, 20, 18, 18]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return workbook, sheet


def export_excel(db_path):
    conn = sqlite3.connect(db_path, timeout=30)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            """
            SELECT ip_address, latitude, longitude, province, city,
                   antenna_gain, polarization, capacity
            FROM devices
            ORDER BY ip_address
            """
        ).fetchall()
    finally:
        conn.close()

    workbook, sheet = _new_workbook()
    for row in rows:
        sheet.append([
            row[field] if row[field] is not None else ""
            for _, field in COLUMNS
        ])
    sheet.auto_filter.ref = sheet.dimensions

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def template_excel():
    workbook, sheet = _new_workbook()
    sheet.append([
        "192.168.1.25",
        "35.6892",
        "51.3890",
        "Tehran",
        "Tehran",
        "30 dBi",
        "H",
        "500 Mbps",
    ])
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
